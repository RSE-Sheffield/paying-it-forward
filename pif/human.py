"""Human reviewers: export a review sheet, import it back as a run.

`pif human export` writes an .xlsx for one reviewer with one row per paper,
a category drop-down, columns for reasoning and minutes spent, the task and
criteria they are classifying against, and a hidden sheet naming the
experiment (with its factor fingerprints). `pif human import` reads a
completed sheet back, checks it, and stores it as the next run of that
experiment, with the sheet itself kept beside the judgements. Importing a
second sheet for the same experiment is a repeat pass.
"""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from . import registry
from .registry import KINDS
from .store import Experiment, ExperimentKey, code_version, now
from .taxonomy import CATEGORIES, VALID, label

REVIEW_SHEET, INFO_SHEET, META_SHEET = "Review", "Instructions", "_experiment"
SHEET_VERSION = 1
PAPER_COLUMNS = ["paper_id", "title", "journal", "year", "field", "doi"]
ENTRY_COLUMNS = ["category", "reasoning", "minutes"]


class SheetError(Exception):
    pass


def _check_human(key: ExperimentKey):
    reviewer = registry.reviewer(key.reviewer)
    if not reviewer.is_human:
        raise SheetError(f"'{key.reviewer}' is a model reviewer; use `pif run` for it")
    return reviewer


def export_sheet(key: ExperimentKey, out: Path) -> Path:
    reviewer = _check_human(key)
    dataset = registry.dataset(key.dataset)
    criteria = registry.criteria(key.criteria)
    prompt = registry.prompt(key.prompt)
    input = registry.input(key.input)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = REVIEW_SHEET
    header = PAPER_COLUMNS + ENTRY_COLUMNS
    ws.append(header)
    for p in dataset.papers:
        ws.append([p.get(c, "") for c in PAPER_COLUMNS] + [None, None, None])
    bold, entry_fill = Font(bold=True), PatternFill("solid", fgColor="FFF5D6")
    for col, name in enumerate(header, 1):
        ws.cell(1, col).font = bold
        if name in ENTRY_COLUMNS:
            for row in range(1, len(dataset.papers) + 2):
                ws.cell(row, col).fill = entry_fill
    widths = {"paper_id": 14, "title": 60, "journal": 28, "year": 6, "field": 18, "doi": 24,
              "category": 10, "reasoning": 70, "minutes": 9}
    for col, name in enumerate(header, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = widths[name]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    cat_col = openpyxl.utils.get_column_letter(header.index("category") + 1)
    min_col = openpyxl.utils.get_column_letter(header.index("minutes") + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(CATEGORIES) + '"', allow_blank=True)
    dv.add(f"{cat_col}2:{cat_col}{len(dataset.papers) + 1}")
    mins = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    mins.add(f"{min_col}2:{min_col}{len(dataset.papers) + 1}")
    ws.add_data_validation(dv)
    ws.add_data_validation(mins)

    info = wb.create_sheet(INFO_SHEET)
    lines = [
        f"Review sheet for {reviewer.name}",
        f"Dataset: {dataset.id} — {dataset.spec.get('title', '')}",
        f"Criteria version: {criteria.id}    Instructions version: {prompt.id}    What to read: {input.id} — {input.description}",
        "",
        "On the Review sheet, fill in the shaded columns for every paper:",
        "  category — one of the nine codes (drop-down)",
        "  reasoning — a short justification",
        "  minutes — roughly how long the paper took you (optional, but used to compare effort with the models)",
        "Don't edit the other columns or the hidden _experiment sheet; they identify the experiment on import.",
        "",
        "Categories:",
        *[f"  {c}  {label(c)}" for c in CATEGORIES],
        "",
        "=== TASK ===", *prompt.task.splitlines(),
        "", "=== CRITERIA ===", *criteria.text.splitlines(),
    ]
    for line in lines:
        info.append([line])
    info.column_dimensions["A"].width = 140
    for row in info.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    meta = wb.create_sheet(META_SHEET)
    fingerprints = Experiment.current_fingerprints(key)
    meta.append(["sheet_version", SHEET_VERSION])
    meta.append(["exported_at", now()])
    for k in KINDS:
        meta.append([k, getattr(key, k)])
    for k in KINDS:
        meta.append([f"fingerprint.{k}", fingerprints[k]])
    meta.sheet_state = "hidden"

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def _text(value) -> str:
    """Cell -> string; Excel stores ids and years as floats (38110704.0)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_sheet_meta(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if META_SHEET not in wb.sheetnames:
        return {}
    return {_text(r[0]): _text(r[1]) for r in wb[META_SHEET].iter_rows(values_only=True) if r and r[0]}


def import_sheet(path: Path, key: ExperimentKey | None = None, *, allow_incomplete: bool = False,
                 notes: str | None = None, reviewed_on: str | None = None,
                 root: Path | None = None, dry_run: bool = False) -> tuple[ExperimentKey, list[dict], str | None]:
    """Read a completed review sheet and store it as a new run.

    The experiment comes from the sheet's hidden _experiment sheet; pass `key`
    for a sheet made some other way (any sheet with an id column, a category
    or classification column, and optionally reasoning and minutes).
    Returns (key, judgements, run id); the run id is None on a dry run.
    """
    meta = read_sheet_meta(path)
    if key is None:
        if not meta:
            raise SheetError(f"{path} has no {META_SHEET} sheet; say which experiment it belongs to "
                             "with --dataset/--criteria/--prompt/--reviewer (and --input if not fulltext)")
        key = ExperimentKey(**{k: meta[k] for k in KINDS})
    _check_human(key)
    dataset = registry.dataset(key.dataset)

    if meta:
        current = Experiment.current_fingerprints(key)
        changed = [k for k in KINDS if meta.get(f"fingerprint.{k}") and meta[f"fingerprint.{k}"] != current[k]]
        if changed:
            raise SheetError(f"{path}: {', '.join(changed)} changed since this sheet was exported, "
                             "so the reviewer worked from different definitions than the registry now has")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[REVIEW_SHEET] if REVIEW_SHEET in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = [_text(h).lower() for h in rows[0]]

    def col(*names):
        return next((header.index(n) for n in names if n in header), None)

    id_col = col("paper_id", str(dataset.spec.get("paper_id", "")).lower())
    cat_col = col("category", "classification")
    reason_col, min_col = col("reasoning"), col("minutes")
    if id_col is None or cat_col is None:
        raise SheetError(f"{path}: needs a paper id column (paper_id or {dataset.spec.get('paper_id')}) "
                         f"and a category column; found {header}")

    order = {p["paper_id"]: i for i, p in enumerate(dataset.papers)}
    judgements, problems, seen = [], [], set()
    for n, row in enumerate(rows[1:], start=2):
        pid = _text(row[id_col])
        if not pid:
            continue
        cat = _text(row[cat_col]).upper()
        if pid not in order:
            problems.append(f"row {n}: paper {pid} is not in dataset {dataset.id}")
        elif pid in seen:
            problems.append(f"row {n}: paper {pid} appears twice")
        elif cat and cat not in VALID:
            problems.append(f"row {n}: '{cat}' is not one of the nine category codes")
        seen.add(pid)
        if not cat:
            continue
        j = {"paper_id": pid, "category": cat,
             "reasoning": _text(row[reason_col]) if reason_col is not None else ""}
        if min_col is not None and _text(row[min_col]):
            try:
                j["minutes"] = float(row[min_col])
            except (TypeError, ValueError):
                problems.append(f"row {n}: minutes '{row[min_col]}' is not a number")
        judgements.append(j)
    missing = [p for p in order if p not in {j["paper_id"] for j in judgements}]
    if missing and not allow_incomplete:
        problems.append(f"{len(missing)} paper(s) have no category: {' '.join(missing[:10])}"
                        + (" ..." if len(missing) > 10 else "") + " (pass --allow-incomplete to import anyway)")
    if problems:
        raise SheetError(f"{path}:\n  " + "\n  ".join(problems))
    judgements.sort(key=lambda j: order[j["paper_id"]])

    if dry_run:
        return key, judgements, None
    exp = Experiment.open(key, root)
    run = exp.new_run({
        "status": "complete" if not missing else "incomplete",
        "started_at": meta.get("exported_at") or None,
        "finished_at": None,
        "imported_at": now(),
        "reviewed_on": reviewed_on,
        "papers_total": len(dataset.papers),
        "papers_judged": len(judgements),
        "failures": [],
        "notes": notes,
        "code": code_version(),
    })
    run.write_judgements(judgements)
    dest = run.dir / f"source{path.suffix}"
    shutil.copy2(path, dest)
    run.meta["source_file"] = dest.name
    run.save()
    return key, judgements, run.id
