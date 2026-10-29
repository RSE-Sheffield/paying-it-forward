"""Tests for the experiment pipeline, against a tiny registry built in a temp dir.

    uv run pytest
"""
import json
from dataclasses import replace

import openpyxl
import pytest

import requests

from pif import cli, human, metrics, paths, prompting, providers, registry, runner
from pif.store import Experiment, ExperimentKey, FactorChanged

KEY = ExperimentKey(dataset="tiny", criteria="v1", prompt="v1", input="fulltext", reviewer="fake")
HUMAN_KEY = ExperimentKey(dataset="tiny", criteria="v1", prompt="v1", input="fulltext", reviewer="human-a")
NIMBLE_KEY = ExperimentKey(dataset="tiny", criteria="v1", prompt="v1", input="abstract-only", reviewer="nimble")


@pytest.fixture
def repo(tmp_path, monkeypatch):
    """A three-paper registry and an empty experiments/ under tmp_path."""
    for name in ("DATASETS", "CRITERIA", "PROMPTS", "INPUTS", "REVIEWERS", "EXPERIMENTS"):
        monkeypatch.setattr(paths, name, tmp_path / name.lower())
    monkeypatch.setattr(paths, "ROOT", tmp_path)
    ds = tmp_path / "datasets" / "tiny"
    (ds / "content").mkdir(parents=True)
    (ds / "dataset.yaml").write_text(
        "title: Tiny\ncontent:\n  text:\n    format: jsonl\n    path: content/text.jsonl\n")
    (ds / "papers.csv").write_text("paper_id,title,journal,year\np1,One,J,2020\np2,Two,J,2021\np3,Three,J,2022\n")
    (ds / "content" / "text.jsonl").write_text("".join(
        json.dumps({"paper_id": p, "abstract": f"abstract {p}", "text": f"body {p} " * 50}) + "\n"
        for p in ("p1", "p2", "p3")))
    (tmp_path / "criteria" / "v1").mkdir(parents=True)
    (tmp_path / "criteria" / "v1" / "criteria.md").write_text("S1 means generated.")
    (tmp_path / "prompts" / "v1").mkdir(parents=True)
    (tmp_path / "prompts" / "v1" / "system.md").write_text("TASK: {task}\nCRITERIA: {criteria}\n")
    (tmp_path / "prompts" / "v1" / "task.md").write_text("Classify it.")
    (tmp_path / "inputs").mkdir()
    (tmp_path / "inputs" / "fulltext.yaml").write_text(
        "body: text\ntemplate: |\n  TITLE: {title}\n  ABSTRACT: {abstract}\n  TEXT: {text}\n")
    (tmp_path / "inputs" / "abstract-only.yaml").write_text(
        "body: none\ntemplate: |\n  TITLE: {title}\n  ABSTRACT: {abstract}\n")
    (tmp_path / "reviewers").mkdir()
    (tmp_path / "reviewers" / "fake.yaml").write_text(
        "description: a stub\nkind: model\nprovider: anthropic\nmodel: fake\n")
    (tmp_path / "reviewers" / "human-a.yaml").write_text("kind: human\nname: A\n")
    (tmp_path / "reviewers" / "nimble.yaml").write_text(
        "kind: model\nprovider: ollama-systemone\nmodel: nimble:9b\n")
    registry.load.cache_clear()
    yield tmp_path
    registry.load.cache_clear()


class FakeProvider:
    """Answers S1D1, except that it fails the papers in `fail` once each."""

    def __init__(self, fail=()):
        self.fail = set(fail)
        self.seen = []

    def describe(self):
        return {}

    def classify(self, system, message):
        text, _ = message.render()
        pid = text.split("body ")[1].split()[0]
        self.seen.append(pid)
        if pid in self.fail:
            self.fail.discard(pid)
            raise ConnectionError("simulated outage")
        return {"category": "S1D1", "reasoning": "stub", "truncated": False,
                "prompt_tokens": 1, "output_tokens": 1, "duration_s": 0.0}


def test_prompts_render_from_versions(repo):
    system = prompting.system_prompt(registry.prompt("v1"), registry.criteria("v1"))
    assert system == "TASK: Classify it.\nCRITERIA: S1 means generated."
    msg = prompting.user_message(registry.dataset("tiny"), registry.input("fulltext"), registry.dataset("tiny").papers[0])
    text, truncated = msg.render()
    assert text.startswith("TITLE: One\nABSTRACT: abstract p1\nTEXT: body p1") and not truncated
    short, truncated = msg.render(max_chars=10)
    assert truncated and short.endswith(prompting.TRUNCATION_MARK + "\n")


def test_failed_papers_are_retried_on_rerun(repo, monkeypatch):
    fake = FakeProvider(fail={"p2"})
    monkeypatch.setattr(runner.providers, "make", lambda reviewer: fake)
    exp = runner.run_experiment(KEY, repeats=1, log=lambda *a: None)
    [run] = exp.runs()
    assert run.status == "incomplete" and run.judged_ids() == {"p1", "p3"}

    fake.seen.clear()
    runner.run_experiment(KEY, repeats=2, log=lambda *a: None)
    runs = exp.runs()
    assert [r.status for r in runs] == ["complete", "complete"]
    assert fake.seen == ["p2", "p1", "p2", "p3"]  # resume judges only p2, then a fresh run
    assert [j["paper_id"] for j in runs[0].judgements()] == ["p1", "p3", "p2"]


def test_editing_a_used_version_is_refused(repo, monkeypatch):
    monkeypatch.setattr(runner.providers, "make", lambda reviewer: FakeProvider())
    runner.run_experiment(KEY, log=lambda *a: None)

    reviewer = repo / "reviewers" / "fake.yaml"
    reviewer.write_text(reviewer.read_text().replace("a stub", "a stub, reworded"))
    dataset = repo / "datasets" / "tiny" / "dataset.yaml"
    dataset.write_text(dataset.read_text() + "    build: some new rebuild instructions\n")
    registry.load.cache_clear()
    Experiment.open(KEY)  # descriptions and build notes can change

    (repo / "criteria" / "v1" / "criteria.md").write_text("S1 means something else.")
    registry.load.cache_clear()
    with pytest.raises(FactorChanged, match="criteria 'v1'"):
        Experiment.open(KEY)


def test_human_sheet_round_trip(repo):
    sheet = repo / "review.xlsx"
    human.export_sheet(HUMAN_KEY, sheet)
    wb = openpyxl.load_workbook(sheet)
    ws = wb["Review"]
    header = [c.value for c in ws[1]]
    for row, cat in zip(ws.iter_rows(min_row=2), ["s1d1", "S2D2", "S3D3"]):
        row[header.index("category")].value = cat
        row[header.index("minutes")].value = 5
    wb.save(sheet)

    key, judgements, run_id = human.import_sheet(sheet)
    assert key == HUMAN_KEY and run_id == "run-01"
    assert [j["category"] for j in judgements] == ["S1D1", "S2D2", "S3D3"]
    assert judgements[0]["minutes"] == 5
    assert human.import_sheet(sheet)[2] == "run-02"  # a second pass is a repeat run

    ws.cell(2, header.index("category") + 1).value = "S9D9"
    wb.save(sheet)
    with pytest.raises(human.SheetError, match="not one of the nine"):
        human.import_sheet(sheet)


def test_metrics():
    assert metrics.modal(["S2D1", "S1D1", "S2D1", "S1D1"]) == "S1D1"  # ties break in taxonomy order
    assert metrics.is_tied(["S2D1", "S1D1"]) and not metrics.is_tied(["S2D1", "S2D1", "S1D1"])
    a = metrics.agreement({"p1": "S1D1", "p2": "S2D2", "p3": "S3D3"}, {"p1": "S1D1", "p2": "S2D3", "p4": "S1D1"})
    assert a["n"] == 2 and a["exact"] == 0.5 and a["s"] == 1.0 and a["d"] == 0.5
    perfect = metrics.counts_table([["S1D1"] * 3, ["S2D2"] * 3], ["S1D1", "S2D2"])
    assert metrics.fleiss_kappa(perfect) == pytest.approx(1.0)


class FakeSystemOne:
    """Stands in for Ollama's /v1/systemone: S2D1 is most probable, S1D1 and
    S3D3 tie for second. Records each request body."""

    def __init__(self):
        self.requests = []

    def post(self, url, timeout, **kwargs):
        assert url.endswith("/v1/systemone")
        self.requests.append(kwargs["json"])
        probs = {c: 0.01 for c in kwargs["json"]["questions"]["category"]["criteria"]}
        probs.update(S2D1=0.8, S1D1=0.05, S3D3=0.05)
        body = {"answers": {"category": {"type": "choice", "choice": "S2D1",
                                         "probabilities": probs, "confidence": 0.7}},
                "usage": {"input_tokens": 123, "output_tokens": 0}}
        resp = requests.Response()
        resp.status_code, resp._content = 200, json.dumps(body).encode()
        return resp


def offline(*args, **kwargs):
    raise requests.ConnectionError("no Ollama in tests")


def test_system_one_judgement_keeps_probabilities(repo, monkeypatch):
    fake = FakeSystemOne()
    monkeypatch.setattr(providers.requests, "post", fake.post)
    monkeypatch.setattr(providers.requests, "get", offline)  # describe() looks up the model digest
    exp = runner.run_experiment(NIMBLE_KEY, log=lambda *a: None)
    [run] = exp.runs()
    assert run.status == "complete"

    sent = fake.requests[0]
    assert sent["state"] == "TITLE: One\nABSTRACT: abstract p1\n"
    question = sent["questions"]["category"]
    assert question["type"] == "choice" and question["instructions"] == "TASK: Classify it.\nCRITERIA: S1 means generated."
    assert list(question["criteria"]) == ["S1D1", "S2D1", "S3D1", "S1D2", "S2D2", "S3D2", "S1D3", "S2D3", "S3D3"]

    j = run.judgements()[0]
    assert j["category"] == "S2D1" and j["confidence"] == 0.7 and j["probabilities"]["S3D3"] == 0.05
    assert j["reasoning"] == "S2D1 0.800 · S1D1 0.050 · S3D3 0.050 (confidence 0.700)"  # ties in taxonomy order
    assert j["prompt_tokens"] == 123 and j["truncated"] is False


def test_system_one_reads_abstracts_only(repo, capsys):
    with pytest.raises(runner.RunError, match="abstract only"):
        runner.run_experiment(replace(NIMBLE_KEY, input="fulltext"), log=lambda *a: None)

    cli.main(["run", "--dataset", "tiny", "--criteria", "v1", "--prompt", "v1", "--reviewer", "nimble",
              "--input", "fulltext", "abstract-only", "--dry-run", "--experiments-dir", str(repo / "experiments")])
    out = capsys.readouterr().out
    assert "skipping tiny/criteria-v1__prompt-v1__fulltext__nimble: 'nimble' is a System One model" in out
    assert "tiny/criteria-v1__prompt-v1__abstract-only__nimble: 0 run(s) -> start 1" in out

    (repo / "reviewers" / "nimble.yaml").write_text(
        "kind: model\nprovider: ollama-systemone\nmodel: nimble:9b\nsettings:\n  temperature: 0.5\n")
    registry.load.cache_clear()
    with pytest.raises(providers.ProviderError, match="temperature"):  # the endpoint would ignore it
        providers.make(registry.reviewer("nimble"))
